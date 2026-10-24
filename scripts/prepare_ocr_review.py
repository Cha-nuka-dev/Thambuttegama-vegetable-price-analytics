"""Prepare a private human-review workbook from the 2025–2026 OCR spreadsheet.

No OCR row is ever automatically marked as verified.
Usage:
    python -m pip install openpyxl
    python scripts/prepare_ocr_review.py path/to/OCR_Review.xlsx --output .local/verification
    # After humans complete the Review Status cells, rerun:
    python scripts/prepare_ocr_review.py --reviewed .local/verification/ocr_verification_queue.xlsx --output .local/verification
Outputs are private (.local is gitignored). Do not commit source-derived data.
"""
from __future__ import annotations

import argparse
import csv
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

CROP_IDS = {
    "Tomato": "tomato", "Brinjal": "brinjal", "Green chilli": "green-chilli",
    "Okra": "okra", "Bitter gourd": "bitter-gourd",
    "Thumba karawila": "spiny-gourd", "Spiny gourd": "spiny-gourd",
    "Capsicum": "capsicum", "Beans": "beans",
    "Local big onion": "local-big-onion", "Cauliflower": "cauliflower",
}
FIELDS = ["Date", "Crop", "Crop ID", "Low Rs/kg", "High Rs/kg", "Midpoint Rs/kg",
          "OCR quality", "Recognized label", "Photo ID", "Source URL",
          "Review Status", "Review Notes", "Review Flags"]
SOURCE_FIELDS = ["Date", "Crop", "Year", "Month no.", "Wholesale low (Rs/kg)",
                 "Wholesale high (Rs/kg)", "Midpoint (Rs/kg)", "OCR quality",
                 "Table column", "Recognized label", "Photo ID", "Facebook source"]


def date_text(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return str(value or "").split(" ")[0]


def numeric(value):
    try:
        return float(value) if value is not None and str(value).strip() else None
    except (TypeError, ValueError):
        return None


def read_source(path):
    book = load_workbook(path, read_only=True, data_only=True)
    if "Selected Prices" not in book:
        raise ValueError("Missing 'Selected Prices' worksheet")
    rows = book["Selected Prices"].iter_rows(values_only=True)
    header = None
    output = []
    for cells in rows:
        values = list(cells)
        if header is None:
            if "Facebook source" in values and "Crop" in values:
                header = {str(x): i for i, x in enumerate(values) if x is not None}
            continue
        if not values or not values[header["Date"]]:
            continue
        def get(field):
            pos = header.get(field)
            return values[pos] if pos is not None and pos < len(values) else None
        d, crop = date_text(get("Date")), str(get("Crop") or "").strip()
        if not d.startswith(("2025-", "2026-")):
            continue
        low = numeric(get("Wholesale low (Rs/kg)"))
        high = numeric(get("Wholesale high (Rs/kg)"))
        label = str(get("Recognized label") or "")
        flags = []
        if crop not in CROP_IDS:
            flags.append("UNKNOWN_CROP")
        if low is None or high is None or low <= 0 or high <= 0:
            flags.append("MISSING_OR_NONPOSITIVE_PRICE")
        elif low > high:
            flags.append("LOW_EXCEEDS_HIGH")
        if "onion" in crop.lower() and ("ආනයන" in label or "import" in label.lower()):
            flags.append("IMPORTED_ONION_NOT_LOCAL")
        if crop == "Beans" and ("සෝයා" in label or "soy" in label.lower()):
            flags.append("SOY_BEANS_VARIETY")
        midpoint = (low + high) / 2 if low is not None and high is not None else None
        output.append([d, crop, CROP_IDS.get(crop, ""), low, high, midpoint,
                       get("OCR quality"), label, str(get("Photo ID") or ""),
                       get("Facebook source"), "Needs review", "", "; ".join(flags)])
    book.close()
    seen = Counter((r[0], r[2]) for r in output)
    for row in output:
        if seen[(row[0], row[2])] > 1:
            row[12] = "; ".join(filter(None, [row[12], "MULTIPLE_ROWS_SAME_DAY_CROP"]))
    return output


def save_review(rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Review Queue"
    ws.append(FIELDS)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    widths = [15, 19, 20, 15, 15, 18, 18, 31, 24, 65, 19, 48, 43]
    for col, width in enumerate(widths, 1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(col)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.alignment = Alignment(wrap_text=True)
    validation = DataValidation(type="list",
        formula1='"Needs review,Verified,Rejected"', allow_blank=False)
    ws.add_data_validation(validation)
    if ws.max_row > 1:
        validation.add(f"K2:K{ws.max_row}")
    info = wb.create_sheet("Instructions")
    for row in [
        ["Purpose", "Human validation before any public data release."],
        ["Review Status", "All OCR rows start as Needs review. OCR quality is not manual verification."],
        ["Verified", "Set only after checking date, crop, variety and BOTH prices against the original image."],
        ["Rejected", "Use for incorrect classifications, bad prices or unsuitable source records."],
        ["Review Notes", "Document image checks, varieties and corrections; fix numeric/crop cells as needed."],
        ["Review Flags", "Automatic warnings; review even if blank."],
        ["Privacy", "Keep this workbook private; do not add OCR data to public GitHub."],
    ]:
        info.append(row)
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 115
    wb.save(path)


def read_review(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = wb["Review Queue"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    wb.close()
    return [list(row[:len(FIELDS)]) for row in rows if row and row[0]]


def export_csv(rows, output):
    verified = []
    counts = Counter()
    for r in rows:
        status = str(r[10] or "Needs review").strip()
        counts[status] += 1
        low, high = numeric(r[3]), numeric(r[4])
        valid = (str(r[0]).startswith(("2025-", "2026-")) and
                 r[2] in CROP_IDS.values() and low is not None and high is not None
                 and 0 < low <= high and r[9])
        flags = str(r[12] or "")
        # Flagged rows require explicit corrections before inclusion.
        if status == "Verified" and valid and not flags:
            verified.append([r[0], r[2], low, high, r[9], "manually_verified"])
    with (output / "verified_prices.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["date", "crop", "low", "high", "source_url", "verification"])
        w.writerows(verified)
    with (output / "review_summary.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["status", "count"])
        w.writerows(sorted(counts.items()))
        w.writerow(["Eligible verified export", len(verified)])
    print(f"Rows: {len(rows)}; verified/exportable: {len(verified)}; status counts: {dict(counts)}")
    if counts["Verified"] > len(verified):
        print("Some rows marked Verified were excluded because of unresolved flags or invalid fields.")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", nargs="?", type=Path)
    parser.add_argument("--reviewed", type=Path)
    parser.add_argument("--output", type=Path, default=Path(".local/verification"))
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    if args.reviewed:
        rows = read_review(args.reviewed)
    elif args.source:
        rows = read_source(args.source)
        save_review(rows, args.output / "ocr_verification_queue.xlsx")
    else:
        parser.error("Provide a source Excel path or --reviewed workbook.")
    export_csv(rows, args.output)


if __name__ == "__main__":
    main()
